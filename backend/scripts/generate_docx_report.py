import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette
    COLOR_PRIMARY = RGBColor(27, 77, 62)       # Forest Green #1B4D3E
    COLOR_SECONDARY = RGBColor(16, 185, 129)   # Emerald #10B981
    COLOR_DARK = RGBColor(30, 41, 59)          # Slate-800 #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)      # Slate-500 #64748B
    HEX_HEADER_BG = "1B4D3E"
    HEX_ALT_BG = "F8FAFC"
    HEX_BORDER = "CBD5E1"

    # --- Title Banner ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("MangoVision Autonomous Scouting Platform")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("System Implementation & Technical Progress Report")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Box Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Project:", " MangoVision Smart Orchard Platform"), ("Location:", " Salem Heritage Mango Orchard, Tamil Nadu")],
        [("Evaluation Date:", " September 2026"), ("Verification Status:", " 100% Tests Passing • Fully Operational")]
    ]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lbl, val = meta_data[r_idx][c_idx]
            run_lbl = p.add_run(lbl)
            run_lbl.font.bold = True
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.color.rgb = COLOR_DARK
            run_val = p.add_run(val)
            run_val.font.size = Pt(9.5)
            run_val.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Helper for Headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK
        return h

    def add_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.italic = italic
        r.font.color.rgb = COLOR_DARK
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_DARK
        return p

    # --- Section: Executive Summary ---
    add_h1("Executive Summary")
    add_body("MangoVision is an advanced cyber-physical precision agriculture platform engineered to automate disease detection, overhead rail robotic scouting, agronomic diagnosis, and remedial treatment across commercial mango orchards in Salem, Tamil Nadu.")
    add_body("This document details all implemented architecture layers, including the foliar pathology synthesis engine, FastAPI backend extensions, Three.js 3D digital twin, 2D corridor matrix, microclimate telemetry, batch remediation, and full verification test passes.")

    # --- Section 1: Foliar Pathology Simulation Engine ---
    add_h1("1. Foliar Pathology Simulation Engine")
    add_body("An automated scouting camera travels along an overhead steel gantry track inspecting 4 orchard rows with 24 commercial trees. To enable realistic agronomic validation, ML benchmark testing, and operator training, a targeted foliar pathology synthesis engine was developed in backend/app/services/simulation_service.py.")
    add_body("The engine synthesizes authentic visual disease lesions, leaf morphology, computer vision bounding boxes, and diagnostic HUD overlays for all 7 common mango pathogens:")

    # Table of Diseases
    t_diseases = doc.add_table(rows=1, cols=4)
    t_diseases.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Pathology", "Scientific Pathogen", "Visual Symptoms Rendered", "Agronomic Severity"]
    col_widths = [Inches(1.3), Inches(1.8), Inches(2.6), Inches(1.1)]

    hdr_cells = t_diseases.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

    disease_rows = [
        ("Anthracnose", "Colletotrichum gloeosporioides", "Irregular dark necrotic brown/black lesions, concentric acervuli rings, shot-hole tears, leaf blade wither", "CRITICAL / HIGH"),
        ("Powdery Mildew", "Oidium mangiferae", "White/greyish superficial talcum felt-like fungal mycelium colonies radiating across veins", "HIGH"),
        ("Bacterial Canker", "Xanthomonas campestris pv. mangiferaeindicae", "Water-soaked angular lesions with bright chlorotic halos, dark gummy bacterial exudate", "CRITICAL"),
        ("Die Back", "Lasiodiplodia theobromae", "Terminal apical branch browning, distinct V-shaped necrotic blade wedges, dark vascular streaks", "CRITICAL / HIGH"),
        ("Gall Midge", "Procontarinia matteiana", "Raised wart-like blister pustules on upper lamina, central circular insect emergence exit holes", "MEDIUM"),
        ("Sooty Mould", "Meliola mangiferae", "Dense black velvety fungal mycelial crust obscuring photosynthetic leaf area", "MEDIUM"),
        ("Cutting Weevil", "Deporaus marginatus", "Clean transverse geometric incision cuts across upper leaf blade, severed apical tip", "MEDIUM"),
        ("Healthy Canopy", "Mangifera indica", "Deep lustrous emerald green lamina, intact reticulate venation, clear margins", "OPTIMAL")
    ]

    for r_idx, (d_name, sci_name, symptoms, sev) in enumerate(disease_rows):
        row = t_diseases.add_row()
        bg = HEX_ALT_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([d_name, sci_name, symptoms, sev]):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_DARK
                if c_idx == 0:
                    r.font.bold = True
                if c_idx == 1:
                    r.font.italic = True
                if c_idx == 3 and "CRITICAL" in val:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(190, 18, 60)

    # Set column widths
    for row in t_diseases.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- Section 2: Backend Architecture & Database Expansion ---
    add_h1("2. Backend Architecture & Database Expansion")
    add_body("The backend API service is built on FastAPI, SQLAlchemy ORM, and Pydantic v2. The database schema and services were expanded to support full agronomic workflows:")

    add_h2("2.1 Database Models")
    add_bullet(" Tracks real-time pathogen detections, severity levels (CRITICAL, HIGH, MEDIUM, LOW), acknowledgement state, resolution timestamps, and associated tree/camera IDs.", bold_prefix="Alert Model (app/models/alert.py):")
    add_bullet(" Complete audit trail for agronomic sprays and interventions, tracking chemical name, dosage, intervention type (CHEMICAL, ORGANIC, PRUNING, BIOLOGICAL), applicator name, and treatment notes.", bold_prefix="Treatment Model (app/models/treatment.py):")
    add_bullet(" Extended health status enum to HEALTHY, DISEASE_DETECTED, TREATED, and UNKNOWN, with spatial row/column references.", bold_prefix="Tree Model (app/models/tree.py):")

    add_h2("2.2 REST API Endpoints")
    add_bullet(" Injects targeted foliar disease on any tree or current camera checkpoint with custom severity and notes.", bold_prefix="POST /api/cameras/{id}/simulation/simulate-disease:")
    add_bullet(" Logs individual tree spray, transitions tree state to TREATED, and auto-resolves pending alerts.", bold_prefix="POST /api/treatments:")
    add_bullet(" Batch applies curative fungicide (e.g. Copper Oxychloride 50 WP) across all diseased trees in one click and clears all warnings.", bold_prefix="POST /api/treatments/batch:")
    add_bullet(" Real-time health scores, canopy distribution metrics, and classified pathogen distributions.", bold_prefix="GET /api/analytics/dashboard:")
    add_bullet(" Downloads comprehensive tabular audit reports of all trees, pathogens, and spray logs.", bold_prefix="GET /api/analytics/export/csv:")

    # --- Section 3: Frontend Dashboard & Interactive Features ---
    add_h1("3. Frontend Dashboard & User Interface")
    add_body("The frontend application is built with React, TypeScript, Three.js (React Three Fiber), and Tailwind CSS, providing orchard managers with rich operational visibility.")

    add_h2("3.1 3D Live Digital Twin & 2D Overhead Grid")
    add_bullet(" Realistic 3D rendered orchard parcel showing overhead rail tracks, animated camera carriage, 24 individual mango trees with dynamic health shaders (Green = Healthy, Red = Diseased with pulse ring, Teal = Treated, Amber = Active Scan), and follow/orbit camera controls.", bold_prefix="3D Digital Twin (FarmScene3D.tsx):")
    add_bullet(" High-density matrix showing row corridors, column tags, variety names, and live carriage position. Clicking any tree focuses it immediately in the inspection viewfinder.", bold_prefix="2D Overhead Matrix (LiveFarmMap.tsx):")

    add_h2("3.2 Real High-Resolution Specimen Viewfinder & Lightbox")
    add_bullet(" Replaced placeholder emojis with real camera-captured foliar images loaded directly from /storage/images/.", bold_prefix="Real Leaf Imagery:")
    add_bullet(" Orchardists can click any leaf thumbnail to open a high-resolution specimen lightbox inspecting lesion edges, acervuli rings, and ML diagnostic metrics.", bold_prefix="HD Specimen Lightbox Modal:")

    add_h2("3.3 Agronomic Remediation & Remission Lifecycle")
    add_bullet(" 1-click action on diseased trees to apply prescribed fungicide and transition tree status to TREATED.", bold_prefix="⚡ Apply Remediation & Mark Treated:")
    add_bullet(" 1-click action on treated trees to re-scan foliar condition and verify that lesions have cured, restoring the tree to HEALTHY.", bold_prefix="🌿 Verify Remission & Mark Healthy:")
    add_bullet(" Top-bar action allowing operators to spray all detected foliar lesions orchard-wide in one click.", bold_prefix="⚡ Spray All Infected Trees (Batch Action):")

    add_h2("3.4 Orchard Microclimate & Foliar Infection Risk Bar")
    add_body("Integrated real-time weather and canopy microclimate telemetry:")
    add_bullet("28.4°C (Optimum vegetative growth window)")
    add_bullet("78% (Elevated — triggers Colletotrichum spore release above 75% threshold)")
    add_bullet("64% Saturation")
    add_bullet("7.2 km/h NW (Safe drift limit)")
    add_bullet("HIGH (Fungal spore pressure active)")
    add_bullet("Next 3h 40m Safe (Prior to evening humidity spike)")

    add_h2("3.5 Printable Field Spray Prescription Work Order")
    add_body("Operators can open and print (window.print()) a complete field spray work order listing:")
    add_bullet("Target tree identifiers, row numbers, and column indices.")
    add_bullet("Target foliar pathogen names and severity levels.")
    add_bullet("Prescribed chemicals (e.g. Copper Oxychloride 50 WP at 3.0 g/L, Wettable Sulphur 80 WP at 2.0 g/L).")
    add_bullet("Central Insecticide Board & Registration Committee (CIB&RC) safety protocols: N95 respirator, nitrile gloves, and 14-day Pre-Harvest Interval (PHI).")

    add_h2("3.6 Global Search & Interactive Filtering")
    add_bullet(" Global search bar with shortcut ⌘K / Ctrl+K, quick filter chips (Diseased, Anthracnose, Treated, Healthy), and instant tree jump.", bold_prefix="Quick Search:")
    add_bullet(" Clicking metric cards or pathogen breakdown rows immediately filters the map and triages infected trees.", bold_prefix="Triage Metric Cards:")

    # --- Section 4: Verification & Testing Results ---
    add_h1("4. Verification & Testing Results")
    add_body("The complete platform was verified through automated test suites and live API execution.")

    add_h2("4.1 Automated Backend Test Suite (Pytest)")
    add_body("All 8 integration tests passed with 100% success rate:")

    # Pytest output box
    box_p = doc.add_paragraph()
    box_p.paragraph_format.space_before = Pt(4)
    box_p.paragraph_format.space_after = Pt(8)
    run_code = box_p.add_run(
        "tests/test_auth.py::test_health_check PASSED [12%]\n"
        "tests/test_auth.py::test_register_and_login PASSED [25%]\n"
        "tests/test_farms.py::test_list_and_create_farms PASSED [37%]\n"
        "tests/test_farms.py::test_farm_layout PASSED [50%]\n"
        "tests/test_ml_inference.py::test_ml_mock_prediction PASSED [62%]\n"
        "tests/test_simulation.py::test_camera_and_simulation_step PASSED [75%]\n"
        "tests/test_simulation.py::test_simulate_common_diseases PASSED [87%]\n"
        "tests/test_treatments_and_alerts.py::test_treatments_and_alerts_lifecycle PASSED [100%]\n"
        "\n== 8 passed in 2.85s =="
    )
    run_code.font.name = "Consolas"
    run_code.font.size = Pt(8.5)
    run_code.font.color.rgb = COLOR_DARK

    add_h2("4.2 Frontend Production Bundle Verification")
    add_body("TypeScript compilation and Vite production packaging completed with 0 errors:")
    box_p2 = doc.add_paragraph()
    box_p2.paragraph_format.space_before = Pt(4)
    box_p2.paragraph_format.space_after = Pt(8)
    run_code2 = box_p2.add_run(
        "> tsc && vite build\n"
        "✓ 2098 modules transformed.\n"
        "dist/index.html                     1.05 kB\n"
        "dist/assets/index-BicBM56u.css     46.25 kB\n"
        "dist/assets/index-COLpHpuI.js   1,322.86 kB\n"
        "✓ built in 39.68s (0 errors)"
    )
    run_code2.font.name = "Consolas"
    run_code2.font.size = Pt(8.5)
    run_code2.font.color.rgb = COLOR_DARK

    add_h2("4.3 End-to-End Live System Execution")
    add_bullet(" Authenticated admin@mangovision.com and acquired JWT token.", bold_prefix="Authentication:")
    add_bullet(" Successfully executed batch spray for 9 infected trees; verified diseased count dropped from 9 to 0 and all active alerts auto-resolved.", bold_prefix="Batch Remediation:")
    add_bullet(" Simulated Anthracnose on T-R01-C01 (97.8% ML confidence, generated CRITICAL alert and high-resolution leaf photo).", bold_prefix="Disease Simulation:")
    add_bullet(" Re-scanned T-R01-C01 with healthy foliar signature, verifying full recovery back to HEALTHY.", bold_prefix="Remission Verification:")

    # Save to file
    output_path = r"c:\Users\Keethana.Rajendran\Downloads\MD\Mango\docs\MangoVision_System_Implementation_Report.docx"
    doc.save(output_path)
    print(f"Successfully generated DOCX report at: {output_path}")

    # Also save to artifact directory
    artifact_path = r"C:\Users\Keethana.Rajendran\.gemini\antigravity-ide\brain\bd5c8a54-5570-4df1-9a09-377332f21e8b\MangoVision_System_Implementation_Report.docx"
    doc.save(artifact_path)
    print(f"Successfully saved artifact DOCX report at: {artifact_path}")

if __name__ == "__main__":
    create_report()
